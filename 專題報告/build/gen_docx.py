# -*- coding: utf-8 -*-
"""產生 Word 正式報告 .docx（deliverable #1）"""
import sys
sys.path.insert(0, r"D:\Things\專題報告\build")
import report_content as rc
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\Things\專題報告\deliverables\專題報告_室內無人機自主導航.docx"
if len(sys.argv) > 1:
    OUT = sys.argv[1]
EA = "Microsoft JhengHei"          # 微軟正黑體
EA2 = "Microsoft JhengHei"
LATIN = "Times New Roman"


def set_ea(style_or_run_node, rpr, name):
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), LATIN)
    rfonts.set(qn('w:hAnsi'), LATIN)
    rfonts.set(qn('w:eastAsia'), name)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "（請於 Word 中更新功能變數以產生目錄）"
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(b); r.append(it); r.append(s); r.append(t); r.append(e)
    return run


def rich(paragraph, text, size=None, bold_all=False, italic=False, color=None):
    """Parse **bold** inline markers into runs."""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        run = paragraph.add_run(seg)
        run.font.name = LATIN
        set_ea(None, run._element.get_or_add_rPr(), EA)
        if size:
            run.font.size = Pt(size)
        is_bold = bold_all or (i % 2 == 1)
        run.bold = is_bold
        run.italic = italic
        if color:
            run.font.color.rgb = color
    return paragraph


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_text(cell, text, size=10.5, bold=False, center=False, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, text, size=size, bold_all=bold)
    if white:
        for run in p.runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_table(doc, headers, rows, size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell_text(t.rows[0].cells[j], h, size=size, bold=True, center=True, white=True)
        shade(t.rows[0].cells[j], "4472C4")
    for r in rows:
        cells = t.add_row().cells
        for j, val in enumerate(r):
            cell_text(cells[j], str(val), size=size, center=(j == 0 and len(val) < 14))
    return t


def ensure_TOC_style(doc):
    pass


def main():
    doc = Document()

    # ---- global page & normal style ----
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(3.0)

    normal = doc.styles['Normal']
    normal.font.name = LATIN
    normal.font.size = Pt(12)
    set_ea(None, normal.element.get_or_add_rPr(), EA)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    # heading styles
    def style_h1():
        st = doc.styles['Heading 1']
        st.font.name = LATIN; st.font.size = Pt(18); st.font.bold = True
        st.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_ea(None, st.element.get_or_add_rPr(), EA)
        st.paragraph_format.page_break_before = False
        st.paragraph_format.space_before = Pt(24)
        st.paragraph_format.space_after = Pt(12)
        st.paragraph_format.keep_with_next = True

    def style_h2():
        st = doc.styles['Heading 2']
        st.font.name = LATIN; st.font.size = Pt(15); st.font.bold = True
        st.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_ea(None, st.element.get_or_add_rPr(), EA)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.keep_with_next = True

    def style_h3():
        st = doc.styles['Heading 3']
        st.font.name = LATIN; st.font.size = Pt(13); st.font.bold = True
        st.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_ea(None, st.element.get_or_add_rPr(), EA)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    style_h1(); style_h2(); style_h3()

    # ---- footer page number ----
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    run.font.size = Pt(9)
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = "PAGE"
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    rr = run._r
    rr.append(b); rr.append(it); rr.append(e)

    # ---- cover ----
    for _ in range(4):
        p = doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, rc.COVER_MAIN, size=26, bold_all=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, rc.COVER_SUB, size=15)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, "專題報告", size=16)
    doc.add_paragraph()
    for label, val in rc.COVER_FIELDS:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        rich(p, f"{label}：{val}", size=13)

    # ---- 摘要 (page break before via paragraph) ----
    h = doc.add_heading("摘要", level=1)
    h.paragraph_format.page_break_before = True
    for b in rc.ABSTRACT:
        if b[0] == "p":
            rich(doc.add_paragraph(), b[1])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    rich(p, f"關鍵詞：{rc.KEYWORDS}", bold_all=True)

    # ---- 目錄 (page break before via paragraph) ----
    h = doc.add_heading("目錄", level=1)
    h.paragraph_format.page_break_before = True
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')

    # ---- body ----
    def blocks(bks, base_size=12, sub=None, sec_no=None, table_counter=None, prefix=""):
        pre_idx = 0
        for b in bks:
            t = b[0]
            if t == "p":
                rich(doc.add_paragraph(), b[1], size=base_size)
            elif t == "bullet":
                p = doc.add_paragraph(style=None)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.3)
                rich(p, "•\t" + b[1], size=base_size)
            elif t == "num":
                pre_idx += 1
                p = doc.add_paragraph(style=None)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.3)
                rich(p, f"{pre_idx}.\t" + b[1], size=base_size)
            elif t == "note":
                p = doc.add_paragraph(style=None)
                p.paragraph_format.left_indent = Cm(0.6)
                rich(p, b[1], size=11, italic=False)
                pPr = p._p.get_or_add_pPr()
                pBdr = OxmlElement('w:pBdr')
                l = OxmlElement('w:left'); l.set(qn('w:val'), 'single'); l.set(qn('w:sz'), '12')
                l.set(qn('w:color'), '4472C4')
                pBdr.append(l)
                pPr.append(pBdr)
            elif t == "pre":
                p = doc.add_paragraph(style=None)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(b[1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.find(qn('w:rFonts'))
                if rfonts is None:
                    rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
                rfonts.set(qn('w:ascii'), 'Consolas'); rfonts.set(qn('w:hAnsi'), 'Consolas')
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'F2F2F2')
                pPr.append(shd)
            elif t == "table":
                if len(b) > 3 and b[3]:
                    cap = doc.add_paragraph()
                    cap.paragraph_format.space_before = Pt(6)
                    cap.paragraph_format.space_after = Pt(2)
                    rich(cap, b[3], size=10.5, bold_all=True)
                add_table(doc, b[1], b[2])
            elif t == "img":
                from docx.shared import Cm as _Cm
                path, caption = b[1], b[2] if len(b) > 2 else ""
                p_pic = doc.add_paragraph()
                p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_pic.paragraph_format.space_before = Pt(6)
                p_pic.add_run().add_picture(path, width=_Cm(14.0))
                if caption:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_after = Pt(10)
                    rich(p_cap, caption, size=10.5, bold_all=True)

    counters = {"t1_1": "1-1", "t3_1": "3-1", "t3_2": "3-2", "tA_1": "A-1"}

    first_section = True
    for s in rc.SECTIONS:
        h = doc.add_heading(f"{s['no']}、{s['title']}", level=1)
        if first_section:
            h.paragraph_format.page_break_before = True
            first_section = False
        for sub in s["subs"]:
            doc.add_heading(f"{sub['no']}　{sub['title']}", level=2)
            blocks(sub["blocks"])

    # ---- appendix ----
    h = doc.add_heading(rc.APPENDIX_TITLE, level=1)
    h.paragraph_format.page_break_before = True
    cap = doc.add_paragraph(); cap.paragraph_format.space_after = Pt(4)
    rich(cap, "表 A-1：專題關鍵技術名詞對照", size=10.5, bold_all=True)
    add_table(doc, rc.APPENDIX_TABLE["headers"], rc.APPENDIX_TABLE["rows"])

    # ---- references ----
    h = doc.add_heading("參考文獻", level=1)
    h.paragraph_format.page_break_before = True
    for i, r in enumerate(rc.REFERENCES, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        rich(p, f"[{i}] {r}", size=11)

    # core properties
    cp = doc.core_properties
    cp.title = rc.TITLE
    cp.subject = rc.SUBTITLE
    cp.author = "專題研究團隊"
    cp.keywords = rc.KEYWORDS

    doc.save(OUT)
    print("DOCX_WRITTEN_OK")
    print("PARAS", len(doc.paragraphs), "TABLES", len(doc.tables))


if __name__ == "__main__":
    main()