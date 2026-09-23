# -*- coding: utf-8 -*-
"""docx 完整度驗證（obs 0043 completeness checklist; obs 0039 round-trip）"""
import docx
import re
import sys

DOC = r"D:\Things\專題報告\deliverables\專題報告_室內無人機自主導航.docx"
LOG = r"D:\Things\專題報告\build\check_docx.log"
if len(sys.argv) > 1:
    DOC = sys.argv[1]

d = docx.Document(DOC)
paras = [p.text for p in d.paragraphs]
alltext = "\n".join(paras)
tables = d.tables
for tb in tables:
    for row in tb.rows:
        for c in row.cells:
            for p in c.paragraphs:
                alltext += "\n" + p.text
xml = d.element.xml
instr = "".join(re.findall(r'<w:instrText[^>]*>(.*?)</w:instrText>', xml, re.S))

checks = [
    ("封面主標題", "室內無人機自主導航之感知定位子系統" in alltext),
    ("封面副標題", "純視覺里程計與 Bundle Adjustment" in alltext),
    ("封面次標題 專題報告", "專題報告" in alltext),
    ("封面佔位欄位（待填）", "【請填寫】" in alltext),
    ("摘要", "摘要" in paras),
    ("摘要數據 17/17", "17/17" in alltext),
    ("摘要數據 2.4–2.8 毫米", "2.4–2.8 毫米" in alltext),
    ("摘要數據 154 毫米", "154 毫米" in alltext),
    ("關鍵詞", "關鍵詞" in alltext),
    ("目錄標題", "目錄" in paras),
    ("TOC 功能變數", "TOC \\o \"1-3\"" in instr),
    ("章節一 研究背景與動機", "一、研究背景與動機" in alltext),
    ("章節二 研究範圍界定", "二、研究範圍界定" in alltext),
    ("章節三 系統架構總覽", "三、系統架構總覽" in alltext),
    ("章節四 核心技術方法", "四、核心技術方法" in alltext),
    ("章節五 實驗結果", "五、實驗結果" in alltext),
    ("章節六 延伸模組", "六、延伸模組" in alltext),
    ("章節七 視覺化與導航模擬", "七、視覺化與導航模擬展示" in alltext),
    ("章節八 研究貢獻總結", "八、研究貢獻總結" in alltext),
    ("章節九 限制與未來工作", "九、限制與未來工作" in alltext),
    ("章節十 開源實作參考", "十、開源實作參考" in alltext),
    ("章節十一 結論", "十一、結論" in alltext),
    ("3.1.1 PnP 架構", "增量式 PnP" in alltext),
    ("Bootstrap 死鎖", "Bootstrap" in alltext),
    ("BA 100 倍改善", "3577 像素" in alltext and "26 像素" in alltext),
    ("七道驗證閘門", "P1–P7" in alltext),
    ("合成 ATE", "2.4–2.8 毫米" in alltext),
    ("回歸 26/26", "26/26" in alltext),
    ("真機 seq_real_14", "seq_real_14" in alltext),
    ("棋盤格 108/201", "108/201" in alltext),
    ("431 個三維地圖點", "431" in alltext),
    ("26 個關鍵幀", "26 個關鍵幀" in alltext),
    ("Pre/Post BA", "2.05 像素" in alltext and "1.24 像素" in alltext),
    ("RMS 0.87", "0.87 像素" in alltext),
    ("P1-P7 結果", "P3（直線度）" in alltext),
    ("76.9%", "76.9%" in alltext),
    ("60 倍落差", "60 倍" in alltext),
    ("敏感度表 cy=30.6", "cy=30.6" in alltext),
    ("敏感度表 154.19", "154.19 mm" in alltext),
    ("敏感度表 4.5%", "4.5%" in alltext and "14%" in alltext),
    ("焦距畸變耦合 k3", "k3 係數異常達到 -1640" in alltext),
    ("Sim3 對齊", "Sim3" in alltext),
    ("標記 0.000 毫米", "0.000 毫米" in alltext),
    ("端到端 52 毫米", "52 毫米" in alltext),
    ("5/5 閘門", "5/5" in alltext),
    ("26/26 回歸", "26/26" in alltext),
    ("敏感度表 focal 假說", "290.81 mm" in alltext and "605.80 mm" in alltext),
    ("敏感度表 fx−z 退化", "fx–z" in alltext or "fx-z" in alltext),
    ("待實體驗證", "尚未完成任何實體硬體驗證" in alltext),
    ("限制 捲尺 ±5–10", "±5–10 毫米" in alltext),
    ("航點導航模擬", "航點導航模擬" in alltext),
    ("7.2 模擬章節", "多樓層救援無人機模擬" in alltext),
    ("7.2 ROS2/Gazebo/PX4", "Gazebo Harmonic" in alltext and "PX4 SITL" in alltext),
    ("7.2 三層樓建築", "20×10 公尺之三層樓" in alltext),
    ("7.2 53 顆 ArUco", "53 顆 ArUco" in alltext),
    ("7.2 46 航點", "46 個航點" in alltext),
    ("7.2 UWB 單元測試", "1e-9" in alltext),
    ("7.2 誠實揭露", "尚待完成" in alltext),
    ("7.3 救災價值", "對救災之潛在助益" in alltext),
    ("7.3 降低人員風險", "降低人員風險" in alltext),
    ("7.3 臨時 GPS", "臨時定位基礎設施" in alltext),
    ("7.3 機器先行人後進", "機器先行、人後進" in alltext),
    ("1.1 救災動機", "社會應用情境" in alltext),
    ("9.4 模擬待辦", "端到端整合實跑與 takeoff/hover" in alltext),
    ("附錄 A", "附錄 A" in alltext),
    ("附錄表格 VSLAM", "VSLAM" in alltext),
    ("參考文獻", "參考文獻" in paras),
    ("參考文獻條目", "[1]" in alltext and "[17]" in alltext),
    ("10.1 開源實作參考", "開源實作參考" in alltext),
    ("開源 ORB-SLAM2", "ORB-SLAM2" in alltext),
    ("開源 ORB-SLAM3", "ORB-SLAM3" in alltext),
    ("開源 VINS-Mono", "VINS-Mono" in alltext),
    ("開源 OpenVINS", "OpenVINS" in alltext),
    ("開源 DSO", "DSO" in alltext),
    ("開源 DBoW2", "DBoW2" in alltext),
    ("開源 AprilTag", "AprilTag" in alltext),
    ("評估工具 evo", "evo" in alltext),
    ("5.4 公開基準對照", "與公開基準之量化對照" in alltext),
    ("基準 EuRoC", "EuRoC" in alltext),
    ("survey 引用", "Sharafutdinov" in alltext),
    ("表 10-1 開源實作參考", "本研究之開源實作參考" in alltext),
    ("表 5-1 基準精度", "公開基準上之代表性定位精度" in alltext),
    ("開源位址 stars", "★10.2k" in alltext),
]

# cover fields order
cvr_ok = True
prev = -1
for label in ["系所班級", "專題名稱", "專題子題", "組員", "指導老師", "日期"]:
    idx = alltext.find(label)
    if idx < 0:
        cvr_ok = False
        break
    if idx <= prev:
        cvr_ok = False
        break
    prev = idx

lines = []
fails = 0
for name, ok in checks:
    lines.append(f"{'PASS' if ok else 'FAIL'}  {name}")
    if not ok:
        fails += 1
lines.append(f"TABLE_COUNT={len(tables)} EXPECTED=6")
lines.append(f"COVER_FIELD_ORDER={'PASS' if cvr_ok else 'FAIL'}")
lines.append(f"TOTAL_FAILS={fails}")

with open(LOG, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("VERIFY_DONE", fails)