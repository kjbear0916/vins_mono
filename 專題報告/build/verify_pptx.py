# -*- coding: utf-8 -*-
"""pptx 內容驗證"""
from pptx import Presentation
from pptx.util import Emu

FILE = r"D:\Things\專題報告\deliverables\專題簡報_感知定位子系統.pptx"
LOG = r"D:\Things\專題報告\build\check_pptx.log"

prs = Presentation(FILE)
lines = []
lines.append(f"SLIDES={len(prs.slides)}")
n_slides = len(prs.slides)
alltext = []
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs)
                if t.strip():
                    texts.append(t)
        if sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        texts.append(t)
        if sh.has_chart:
            texts.append("[chart]")
    alltext.append("\n".join(texts))
    lines.append(f"SLIDE_{i}: {len(texts)} text-elements")
full = "\n".join(alltext)

joins = [
    ("封面主標題", "感知定位子系統"),
    ("第二章大綱", "報告大綱"),
    ("背景轉折", "飛行載具損壞"),
    ("問題本質", "視覺里程計"),
    ("範圍表", "Bundle Adjustment 全局優化"),
    ("架構", "MJPEG"),
    ("VO前端", "solvePnPRansac"),
    ("關鍵問題表", "Bootstrap"),
    ("BA", "100 倍"),
    ("開源對照", "相關開源模型與技術定位"),
    ("開源系統", "ORB-SLAM2／ORB-SLAM3"),
    ("開源 VIO", "OpenVINS／Kimera-VIO"),
    ("公開基準 EuRoC", "EuRoC"),
    ("七道閘門", "P7 靜止段光流檢驗"),
    ("合成結果", "17 / 17"),
    ("真機結果", "seq_real_14"),
    ("落差", "60 倍"),
    ("ArUco", "Sim3 對齊"),
    ("視覺化", "航點導航模擬"),
    ("未來工作", "DBoW2＋g2o／GTSAM"),
    ("貢獻開源定位", "與開源生態／公開基準之對照定位"),
    ("結論", "謝謝聆聽"),
]
for name, s in joins:
    ok = s in full
    lines.append(f"{'PASS' if ok else 'FAIL'}  {name}[{s}]")
    print("CHECK", name, ok)

chart_slide = any("[chart]" in t for t in alltext)
lines.append(f"CHART_PRESENT={'PASS' if chart_slide else 'FAIL'}")
with open(LOG, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("SLIDES", n_slides)